from pathlib import Path
import hashlib, json, re
from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader

root = Path(__file__).resolve().parents[1]
old_path = root / 'history/competition-team-brief-edition/FlightOR_商业计划书_终稿.docx'
docx_path = root / 'FlightOR_商业计划书_终稿.docx'
pdf_path = root / 'FlightOR_商业计划书_终稿.pdf'
source_path = Path(r'C:\Users\VENTSDENYE5\Downloads\FlightOR 商业计划书（正式排版版）.docx')

def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def content(doc, exclude_edited_sections=False):
    result = []
    skip_section = False
    for el in doc.element.body:
        text = ''.join(el.xpath('.//w:t/text()')).strip()
        if not text:
            continue
        if exclude_edited_sections:
            if re.match(r'^[一二三四五六七八九十]+  ', text):
                title = re.sub(r'^[一二三四五六七八九十]+  ', '', text)
                skip_section = title == '核心团队背景'
            if skip_section:
                continue
            text = re.sub(r'^[一二三四五六七八九十]+  ', '', text)
            text = re.sub(r'^(图|表) \d+  ', r'\1  ', text)
            text = text.replace('集中在第十一、十二章', '集中在第十、十一章')
        result.append(text)
    return result

old = Document(old_path)
doc = Document(docx_path)
expected = content(old, True)
actual = content(doc)
comparable = content(doc, True)
pdf = PdfReader(pdf_path)
page_texts = [p.extract_text() or '' for p in pdf.pages]
pdf_text = '\n'.join(page_texts)
all_doc_text = '\n'.join(actual)
main_headings = [p for p in doc.paragraphs if re.match(r'^[一二三四五六七八九十]+  ', p.text)]
team_expected = [
    '核心团队由胡斌、曾郡玫和李昊泽组成，教育背景涉及北京航空航天大学、南京大学和清华大学。胡斌与曾郡玫具有本科及硕士教育经历，并曾在多家企业或业务团队实习。',
    '胡斌',
    '教育背景：本科阶段就读于北京航空航天大学，硕士阶段就读于清华大学。',
    '实习经历：曾在鹰角、百度、高德、可灵等企业或业务团队实习。',
    '曾郡玫',
    '教育背景：本科阶段就读于南京大学，硕士阶段就读于清华大学。',
    '实习经历：曾在网易、面壁智能和微软实习。',
    '李昊泽',
    '教育背景：目前为北京航空航天大学本科在读学生，是 FlightOR 核心团队成员之一。',
]
team_actual = []
in_team = False
for p in doc.paragraphs:
    if p.text == '十三  核心团队背景':
        in_team = True
        continue
    if in_team and re.match(r'^[一二三四五六七八九十]+  ', p.text):
        break
    if in_team and p.text.strip():
        team_actual.append(p.text)
before_first_heading = []
for el in doc.element.body:
    if ''.join(el.xpath('.//w:t/text()')).strip() == '一  项目概述':
        break
    before_first_heading.append(el)
checks = {
    'pages': len(pdf.pages),
    'layout': 'continuous_report',
    'chapter_count': len(main_headings),
    'tables': len(doc.tables),
    'pictures': len(doc.inline_shapes),
    'unchanged_content_outside_team': expected == comparable,
    'competition_main_title_count': sum(p.text == '七  竞争格局' for p in doc.paragraphs),
    'old_competition_chapters_merged': not any(p.text in {'七  竞品分析', '八  核心竞争优势与差异化'} for p in doc.paragraphs),
    'no_kicker_paragraphs': not any('EXECUTIVE SUMMARY' in p.text or 'SEARCH AND TRANSACTION' in p.text for p in doc.paragraphs),
    'team_matches_user_supplied_bios': team_actual == team_expected,
    'cover_has_no_diagram': not any(el.xpath('.//w:drawing') for el in before_first_heading),
    'figure_numbers_sequential': [int(re.match(r'^图 (\d+)', p.text)[1]) for p in doc.paragraphs if re.match(r'^图 \d+', p.text)] == list(range(1,10)),
    'table_numbers_sequential': [int(re.match(r'^表 (\d+)', p.text)[1]) for p in doc.paragraphs if re.match(r'^表 \d+', p.text)] == list(range(1,17)),
    'all_tables_repeat_header': all(t.rows[0]._tr.xpath('./w:trPr/w:tblHeader') for t in doc.tables),
    'all_rows_non_splitting': all(r._tr.xpath('./w:trPr/w:cantSplit') for t in doc.tables for r in t.rows),
    'all_images_alt': all(p._inline.docPr.get('descr') for p in doc.inline_shapes),
    'pdf_has_no_missing_glyph_marker': '\ufffd' not in pdf_text,
    'all_page_numbers_match': all(f'{i+1}/{len(pdf.pages)}' in re.sub(r'\s+', '', t) for i, t in enumerate(page_texts) if i>0),
    'source_unchanged': sha(source_path) == 'f477a1bc8c631dd9ba565d6aeac40f503220e7ddcfba217254fe4bafb5ee9319',
    'docx_sha256': sha(docx_path),
    'pdf_sha256': sha(pdf_path),
    'docx_bytes': docx_path.stat().st_size,
    'pdf_bytes': pdf_path.stat().st_size,
}
if expected != comparable:
    from difflib import unified_diff
    (root/'qa/content-diff.txt').write_text('\n'.join(unified_diff(expected, comparable, fromfile='expected', tofile='actual')), encoding='utf-8')
(root/'qa/final-checks.json').write_text(json.dumps(checks, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps(checks, ensure_ascii=False, indent=2))
assert all(v for v in checks.values() if isinstance(v, bool)), 'Document check failed'
assert len(main_headings) == 15 and len(doc.tables) == 16 and len(doc.inline_shapes) == 11
