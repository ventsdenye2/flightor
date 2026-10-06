from pathlib import Path
from copy import deepcopy
import re, json, hashlib, math
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent
SOURCE = Path(r'C:\Users\VENTSDENYE5\Downloads\FlightOR 商业计划书（正式排版版）.docx')
OUT = ROOT / 'FlightOR_商业计划书_终稿.docx'
ASSETS = ROOT / 'qa' / 'diagrams'
ASSETS.mkdir(parents=True, exist_ok=True)
BLUE='245A96'; INK='202B38'; MUTED='586777'; PALE='F1F5FA'; BORDER='D9D9D9'
FONT='Microsoft YaHei'
source = Document(SOURCE)
original = [p.text for p in source.paragraphs]
doc = Document(SOURCE)
for el in list(doc.element.body):
    if el.tag != qn('w:sectPr'): doc.element.body.remove(el)
s = doc.sections[0]
s.page_width = Cm(21); s.page_height = Cm(29.7)
s.top_margin = Cm(2); s.bottom_margin = Cm(1.9)
s.left_margin = Cm(2.2); s.right_margin = Cm(2.2)
s.header_distance = Cm(0.85); s.footer_distance = Cm(0.85)
s.different_first_page_header_footer = True

def fontstyle(style, size, bold=False, color='000000'):
    style.font.name=FONT; style.font.size=Pt(size); style.font.bold=bold
    style.font.color.rgb=RGBColor.from_string(color)
    rp=style.element.get_or_add_rPr()
    rf=rp.find(qn('w:rFonts'))
    if rf is None: rf=OxmlElement('w:rFonts'); rp.insert(0,rf)
    for key in ('ascii','hAnsi','eastAsia','cs'): rf.set(qn('w:'+key),FONT)
    for key in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'): rf.attrib.pop(qn('w:'+key),None)
    for color_el in rp.findall(qn('w:color')):
        for key in ('themeColor','themeTint','themeShade'): color_el.attrib.pop(qn('w:'+key),None)

for st in doc.styles:
    if st.type in (1,2):
        fontstyle(st,11)
        if st.type == 1:
            for e in list(st.element.get_or_add_pPr()):
                if e.tag in (qn('w:pBdr'),qn('w:shd')): e.getparent().remove(e)
normal=doc.styles['Normal']
normal.paragraph_format.line_spacing=Pt(16.5)
normal.paragraph_format.space_after=Pt(8)
normal.paragraph_format.space_before=Pt(0)
normal.paragraph_format.widow_control=True
normal.paragraph_format.first_line_indent=Pt(0)
normal.paragraph_format.keep_with_next=False
normal.paragraph_format.page_break_before=False
sg=OxmlElement('w:snapToGrid');sg.set(qn('w:val'),'0');normal.element.get_or_add_pPr().append(sg)
for name,size,before,after in [('Title',30,0,10),('Subtitle',13,0,10),('Heading 1',15,18,9),('Heading 2',12,12,6),('Heading 3',11,9,6),('Caption',9,6,7)]:
    st=doc.styles[name] if name in doc.styles else doc.styles.add_style(name,1)
    fontstyle(st,size,name.startswith('Heading'))
    pf=st.paragraph_format; pf.space_before=Pt(before);pf.space_after=Pt(after)
    pf.line_spacing=Pt(size*1.35);pf.first_line_indent=Pt(0)
    pf.keep_with_next=name.startswith('Heading');pf.keep_together=True
fontstyle(doc.styles['Caption'],9,color=MUTED)
for name,size,color in [('BP Small',9,MUTED),('BP Table',10.5,INK),('BP Lead',12,INK),('BP Ref',9,INK),('BP Kicker',9.5,MUTED)]:
    st=doc.styles.add_style(name,1) if name not in doc.styles else doc.styles[name]
    st.base_style=normal;fontstyle(st,size,color=color)
    st.paragraph_format.line_spacing=Pt(size*1.4);st.paragraph_format.space_after=Pt(6)
doc.styles['BP Lead'].paragraph_format.space_after=Pt(14)
doc.styles['BP Ref'].paragraph_format.space_after=Pt(8)
doc.styles['BP Table'].paragraph_format.space_after=Pt(0)
doc.styles['BP Table'].paragraph_format.space_before=Pt(0)

def runfont(run,size=None,bold=None,color=None):
    run.font.name=FONT
    for key in ('ascii','hAnsi','eastAsia','cs'): run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:'+key),FONT)
    if size: run.font.size=Pt(size)
    if bold is not None: run.bold=bold
    if color: run.font.color.rgb=RGBColor.from_string(color)

def rich(p,text):
    for i,t in enumerate(text.split('**')):
        r=p.add_run(t);runfont(r,bold=True if i%2 else None)
    return p

def para(text='',style=None):
    p=doc.add_paragraph(style=style or 'Normal');rich(p,text);return p

def heading(text,level=2): return para(text,f'Heading {level}')
sections=[]
def section(title,legacy_kicker=''):
    """Continuous report: heading once per chapter, without per-page display titles."""
    n=len(sections)+1;sections.append({'sequence':n,'title':title})
    if n>1:
        h=heading(title,1)
        # Only the cover has a deliberate page boundary; references also flow.
        h.paragraph_format.page_break_before=n==2
        if h.paragraph_format.page_break_before: h.paragraph_format.space_before=Pt(0)
    return n

def note(text): return para(text,'BP Small')
def lead(text): return para(text,'BP Lead')

tblcount=0; figcount=0
def table(title,headers,rows,widths=None,size=10.5,center_cols=()):
    global tblcount
    tblcount+=1
    cap=para(f'表 {tblcount}  {title}','Caption');cap.paragraph_format.keep_with_next=True
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    widths=widths or [16.6/len(headers)]*len(headers)
    pr=t._tbl.tblPr
    old=pr.find(qn('w:tblStyle'))
    if old is not None: pr.remove(old)
    tw=pr.find(qn('w:tblW'));tw.set(qn('w:w'),str(int(16.6/2.54*1440)));tw.set(qn('w:type'),'dxa')
    borders=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'5');el.set(qn('w:color'),BORDER);borders.append(el)
    pr.append(borders)
    mar=OxmlElement('w:tblCellMar')
    for edge,value in [('top',85),('left',115),('bottom',85),('right',115)]:
        el=OxmlElement('w:'+edge);el.set(qn('w:w'),str(value));el.set(qn('w:type'),'dxa');mar.append(el)
    pr.append(mar)
    for col,w in zip(t.columns,widths): col.width=Cm(w)
    for ri,data in enumerate([headers]+rows):
        row=t.rows[0] if ri==0 else t.add_row()
        trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        if ri==0: trpr.append(OxmlElement('w:tblHeader'))
        for ci,(cell,txt,w) in enumerate(zip(row.cells,data,widths)):
            cell.width=Cm(w);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shading=OxmlElement('w:shd');shading.set(qn('w:fill'),BLUE if ri==0 else (PALE if ri%2==0 else 'FFFFFF'));cell._tc.get_or_add_tcPr().append(shading)
            p=cell.paragraphs[0];p.style=doc.styles['BP Table'];p.paragraph_format.keep_together=True
            # Keep short tables intact; longer tables need at least two data
            # rows together at either end when Word paginates them naturally.
            p.paragraph_format.keep_with_next=(ri < len(rows)) if len(rows)<=4 else (ri<=1 or ri==len(rows)-1)
            if ri==0 or ci in center_cols: p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            rich(p,str(txt))
            for r in p.runs: runfont(r,size=size,bold=True if ri==0 else r.bold,color='FFFFFF' if ri==0 else INK)
    after=doc.add_paragraph();after.paragraph_format.space_after=Pt(2);after.paragraph_format.space_before=Pt(0);after.paragraph_format.line_spacing=1;after.add_run().font.size=Pt(2)
    return t

regular=Path(r'C:\Windows\Fonts\msyh.ttc'); boldfont=Path(r'C:\Windows\Fonts\msyhbd.ttc')
def draw_text_center(draw,xy,text,size,bold=False,fill='#202B38'):
    ft=ImageFont.truetype(str(boldfont if bold else regular),size)
    lines=text.split('\n');line_h=int(size*1.4);total=line_h*len(lines)
    for i,line in enumerate(lines):
        bbox=draw.textbbox((0,0),line,font=ft);w=bbox[2]-bbox[0]
        draw.text((xy[0]-w/2,xy[1]-total/2+i*line_h),line,font=ft,fill=fill)

def flow(name,nodes,cols=4,sub=None):
    W=1600;cols=min(cols,len(nodes));rows=math.ceil(len(nodes)/cols)
    H=210*(rows-1)+184;im=Image.new('RGB',(W,H),'white');d=ImageDraw.Draw(im)
    margin=12;gap=55;bw=(W-margin*2-gap*(cols-1))/cols;bh=138
    centers=[]
    for i,label in enumerate(nodes):
        row=i//cols;col=i%cols
        if row%2: col=cols-1-col
        x=margin+col*(bw+gap);y=22+row*210
        d.rounded_rectangle((x,y,x+bw,y+bh),radius=9,fill='#F1F5FA',outline='#B7C8DC',width=2)
        draw_text_center(d,(x+24,y+20),f'{i+1:02}',18,bold=True,fill='#245A96')
        draw_text_center(d,(x+bw/2,y+bh/2+6),label,35,bold=True)
        centers.append((x+bw/2,y+bh/2,x,y,bw,bh))
    for i in range(len(nodes)-1):
        a=centers[i];b=centers[i+1]
        if i//cols==(i+1)//cols:
            right=b[0]>a[0];x1=a[2]+a[4]+8 if right else a[2]-8;x2=b[2]-12 if right else b[2]+b[4]+12;y=a[1]
            d.line((x1,y,x2,y),fill='#245A96',width=4);sgn=1 if right else -1
            d.polygon([(x2,y),(x2-sgn*12,y-7),(x2-sgn*12,y+7)],fill='#245A96')
        else:
            x=a[0];y1=a[3]+a[5]+7;y2=b[3]-12
            d.line((x,y1,x,y2),fill='#245A96',width=4);d.polygon([(x,y2),(x-7,y2-12),(x+7,y2-12)],fill='#245A96')
    path=ASSETS/(name+'.png');im.save(path,dpi=(245,245));return path

def figure(path,caption,alt,width=16.6):
    global figcount
    figcount+=1;p=doc.add_paragraph();p.paragraph_format.space_after=Pt(1);p.paragraph_format.keep_with_next=True
    p.paragraph_format.line_spacing=1
    r=p.add_run();pic=r.add_picture(str(path),width=Cm(width));pic._inline.docPr.set('descr',alt)
    para(f'图 {figcount}  {caption}','Caption')

def diagram(name,nodes,caption,cols=4):
    figure(flow(name,nodes,cols),caption,' → '.join(n.replace('\n',' ') for n in nodes))

def screenshot_gallery(items,caption,width=4.5):
    global figcount
    figcount+=1
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
    p.paragraph_format.line_spacing=1;p.paragraph_format.space_after=Pt(3)
    for i,(path,alt,crop) in enumerate(items):
        if i: p.add_run('    ')
        im=Image.open(path);cw=im.width*(1-crop[0]-crop[2]);ch=im.height*(1-crop[1]-crop[3])
        pic=p.add_run().add_picture(str(path),width=Cm(width),height=Cm(width*ch/cw))
        pic._inline.docPr.set('descr',alt)
        if any(crop):
            sr=OxmlElement('a:srcRect')
            for attr,val in zip(('l','t','r','b'),crop): sr.set(attr,str(round(val*100000)))
            bf=pic._inline.graphic.graphicData.pic.blipFill;bf.insert(1,sr)
    para(f'图 {figcount}  {caption}','Caption')

def field(p,instr):
    r=p.add_run();b=OxmlElement('w:fldChar');b.set(qn('w:fldCharType'),'begin');r._r.append(b)
    r=p.add_run();it=OxmlElement('w:instrText');it.set(qn('xml:space'),'preserve');it.text=' '+instr+' ';r._r.append(it)
    r=p.add_run();sep=OxmlElement('w:fldChar');sep.set(qn('w:fldCharType'),'separate');r._r.append(sep)
    r=p.add_run('1');runfont(r,9,color=MUTED)
    r=p.add_run();e=OxmlElement('w:fldChar');e.set(qn('w:fldCharType'),'end');r._r.append(e)

for part in (s.header,s.footer,s.first_page_header,s.first_page_footer):
    for el in list(part._element): part._element.remove(el)
for part in (s.first_page_header,s.first_page_footer):
    blank=part.add_paragraph(style=normal)
    blank.paragraph_format.space_after=Pt(0)
    pb=OxmlElement('w:pBdr')
    for side in ('top','bottom','left','right'):
        e=OxmlElement('w:'+side);e.set(qn('w:val'),'nil');pb.append(e)
    blank._p.get_or_add_pPr().append(pb)
hp=s.header.add_paragraph();hp.paragraph_format.space_after=Pt(0)
hp.paragraph_format.tab_stops.add_tab_stop(Cm(16.6),2)
runfont(hp.add_run('FlightOR  商业计划书\t创业孵化'),9,color='000000')
fp=s.footer.add_paragraph();fp.paragraph_format.tab_stops.add_tab_stop(Cm(16.6),2);fp.paragraph_format.space_after=Pt(0)
runfont(fp.add_run('2026.09\t'),9,color=MUTED);field(fp,'PAGE');runfont(fp.add_run(' / '),9,color=MUTED);field(fp,'NUMPAGES')
update=doc.settings.element.find(qn('w:updateFields'))
if update is None: update=OxmlElement('w:updateFields');doc.settings.element.append(update)
update.set(qn('w:val'),'true')

# 01 Cover. Preserve the source's identity and date.
section('封面','')
p=para('','Normal');p.paragraph_format.space_after=Pt(56)
p=para('FlightOR','Title');runfont(p.runs[0],48,bold=True,color='000000')
para('商业计划书','Title')
para('国际自由行航线与路线规划','Subtitle')
p=para('');p.paragraph_format.space_after=Pt(230)
para('创业孵化  产品上线前正式版','Subtitle')
para('版本日期  2026 年 9 月','BP Small')

# 02 Executive summary
section('一  项目概述','项目定位  /  EXECUTIVE SUMMARY')
lead('FlightOR 面向预算和时间有限、目的地或路线尚未完全确定的年轻自由行用户，帮助他们比较国际航线，并形成可持续修改的旅行方案。')
para('产品以微信小程序承载自然语言交互，聚焦从需求理解、航线搜索到路线决策的规划过程。用户先表达预算、可出行时间和兴趣，再查看候选路线、调整条件，并按需生成带来源的当地攻略。')
heading('典型规划场景')
para('用户从北京出发，国庆前后有 7 天时间，预算 1.5 万元，想去日本，喜欢文化和美食，但尚未决定东京、大阪或多城组合。传统方式需要在 OTA、攻略平台、地图和通用 AI 之间反复搬运条件；FlightOR 将需求、候选方案和后续修改组织在同一规划过程内。')
note('场景中的日期、预算和目的地用于说明需求，不代表已完成的用户案例或实时可订报价。')
screenshot_gallery([
    (ROOT.parent/'playwright'/'flight-first'/'live-candidates-11-mobile.png','2026-09-14 H5 真实航班查询的比较区域，历史价格，局部显示',(0,0,0.335,0.327)),
    (ROOT.parent.parent/'docs'/'design'/'budget-travel-agent'/'publication-ui-evidence'/'playwright-selfTicket-zh-overview.png','2026-09-22 正式 H5 行程概览，固定测试行程',(0,0,0,0)),
], 'H5 航班比较与行程概览 左为 9 月 14 日真实查询局部，右为 9 月 22 日固定测试行程',3.7)
table('项目阶段与下一步',['项目维度','目前基础','下一步'],[
['产品形态','已形成可运行小程序、独立后端和 Agent 规划链路','完善正式上线与真实平台体验'],
['工程验证','保留根项目 140/140、后端 121 项历史测试记录','以具体场景和平台验收确认交付质量'],
['市场验证','正式上线前的创业孵化阶段','验证真实需求、有效规划和复访'],
], [2.6,7,7])
para('**核心价值**  在预算、时间和兴趣约束下，降低反复查票与比较路线的成本，帮助用户形成愿意采用的旅行方案。')

# 03 Pain points
section('二  用户痛点与需求背景','用户需求  /  DECISION COST')
lead('自由行规划的核心难点在于信息分散，以及日期、价格、城市顺序和停留时间之间高度关联。')
diagram('pain',['决定目的地','寻找日期与价格','比较交通方案','调整城市与停留'],'旅行决策循环 任一条件变化都可能触发重新比较',4)
para('例如，一张便宜 800 元的机票，如果需要在中转机场停留 14 小时，未必具有更低的综合成本；如果这段时间足以进入市区完成半日游，中转又可能成为旅行体验的一部分。判断需要同时考虑价格、时间、入境与交通条件。')
note('800 元与 14 小时为原稿说明性示例，不构成价格承诺或过境可行性结论。')
table('高摩擦场景与用户要完成的判断',['场景','用户的关键判断'],[
['灵活日期','能否通过前后调整日期，获得更合适的价格和出行时间'],
['目的地未定','哪些目的地同时符合预算、兴趣和可出行时间'],
['多城路线','城市顺序、进出机场和多段交通如何组合'],
['中转利用','节省金额是否值得额外等待，中转时间能否安全利用'],
['多轮修改','改变预算、城市或日期后，能否延续已有规划'],
],[3.2,13.4])
para('OTA 提供库存、交易与售后，航班搜索工具帮助发现价格和日期，内容平台提供经验与灵感，通用 AI 协助表达与整理。FlightOR 将这些信息连接到同一旅行约束下，减少用户跨工具整理和重新比较的工作。')

# 04 Solution
section('三  产品定位与解决方案','产品方案  /  USER JOURNEY')
lead('FlightOR 定位为以航线与约束优化为核心的 AI 自由行规划工具。')
para('规划页作为微信小程序首个入口，用户可直接用自然语言描述需求。系统将对话转化为出发地、日期、天数、预算、区域、兴趣、必去与排除城市等条件，并单独展示“当前理解”，供用户检查与修正。')
diagram('journey',['自然语言需求','结构化旅行条件','目的地与路线探索','航班与路线卡','确认或调整','实时报价核对','按需生成攻略','保存与继续规划'],'产品交互流程示意',4)
heading('结构化状态支持持续修改')
para('FlightOR 将自然语言对话持续映射为可执行、可修改的旅行状态。用户补充条件时，系统更新相应字段，并基于当前状态继续规划。')
para('“预算再加两千”更新预算约束；“不要东京”更新排除城市；“最好周五晚上走”补充出发日期与时间偏好。')
screenshot_gallery([
    (ROOT.parent/'playwright'/'flight-first'/'planner-mobile-390x844.png','2026-09-14 正式 H5 规划入口截图',(0,0,0,0)),
    (ROOT.parent.parent/'docs'/'design'/'budget-travel-agent'/'place-media-evidence'/'takeshita-detail.png','2026-09-22 正式 H5 景点详情，固定测试行程，真实 Wikimedia 照片及页面许可',(0,0,0,0)),
], 'H5 规划入口与景点详情 详情使用固定测试行程及真实授权照片',4.2)
note('截图记录于 2026-09-14、2026-09-22，展示界面与对应测试范围；竹下通照片：Intforce，CC BY-SA 4.0，来源与许可见 [12]。')

# 05 Core capabilities
section('四  核心产品能力与产品架构','产品能力  /  CAPABILITIES')
table('六项能力与实施范围',['能力','用户获得的价值','范围与后续计划'],[
['对话式需求理解','将口语需求转为结构化旅行约束','已建立中文预算确定性解析；“一万五”“1万5”“1.5万”保持一致含义'],
['灵活日期与路线探索','在日期与目的地尚有弹性时比较方案','当前最终航线生成以单出发机场、单最终目的地、单程和有界日期为主'],
['多城路线与中转利用','综合判断城市组合与中转的价格和时间代价','多城、完整往返与地面段组合继续深化；总耗时、过夜、签证、机场切换、行李和游玩时间逐步纳入评分'],
['路线卡与报价确认','比较候选路线并核对价格','规划结果与价格确认分层；采用航班方案不等于锁价或购票'],
['带来源的旅行攻略','按需获得景点和每日玩法','保留网页、目录与规则等来源类型；来源可追溯不等于当前事实全部核实'],
['会话状态与用户记忆','保留历史并延续修改','当前已有会话历史与服务端旅行状态；跨会话偏好在授权和隐私合规前提下继续完善'],
],[3.1,5.7,7.8],10.5)
heading('规划与价格确认分别承担责任')
diagram('fact_layers',['AI 与 Agent\n理解 规划 解释','工具与数据\n查询 价格核对','用户确认\n检查条件与方案'],'生成内容与外部事实分层',3)
para('微信小程序负责交互和展示，自建后端统一管理 Agent、航班搜索、研究与攻略服务。第三方密钥由后端管理，模型与数据服务通过适配层接入，支持按业务需要替换。')
note('实施范围依据项目当前文档 [11]；接口和构建细节集中在第十、十一章，不将后续能力列为已验收成果。')

# 06 Personas
section('五  目标用户与首批用户画像','首批用户  /  TARGET USERS')
lead('第一阶段聚焦 20—35 岁、每年有 1—3 次自主旅行决策、习惯自行查票与规划行程的年轻自由行用户。')
para('用户筛选更关注自主规划意愿、价格或路线优化需求，以及对非标准路线的接受度。以下三类画像用于定义招募方向，属于目标用户描述。')
table('三类首批目标用户',['维度','A 预算敏感的灵活出行者','B 目的地未定的探索者','C 多城自由行用户'],[
['典型情况','学生、初入职场年轻人或数字游民','已有假期和预算，只有区域或兴趣目标','已确定国家或区域，希望一次走 2—4 个城市'],
['可调整条件','出行日期可调整 1—5 天，可考虑中转或邻近机场','目的地与城市组合尚可调整','城市顺序、进出机场与交通组合'],
['核心需求','以可接受的时间代价找到更合适的票价','把兴趣、预算、时间和交通成本合并比较','评估多段交通对整体路线的影响'],
['产品价值','减少重复搜索，发现替代路线','形成符合约束的候选目的地与方案','持续调整完整路线，降低组合比较成本'],
],[2.1,4.85,4.85,4.8],10.5)
heading('共同的行为特征')
para('**自主规划**  愿意参与条件设定与方案选择。')
para('**存在优化空间**  日期、预算、目的地或航线至少有一项可调整。')
para('**接受路线取舍**  愿意比较更低价格与额外时间、中转或机场变化之间的代价。')
note('用户年龄、频次及行为范围为项目定位，非人口比例调查或已验证市场渗透率。')

# 07 Market
section('六  市场分析','市场依据  /  MARKET CONTEXT')
lead('2025 年内地居民出境旅游 1.4836 亿人次，同比增长 20.8%。这一市场基础支持跨境规划需求的存在，但不直接等于 FlightOR 的可触达用户规模。[2]')
table('已核查的旅游市场数据',['统计范围','时间','规模或金额','同比'],[
['内地居民出境旅游','2025 全年','1.4836 亿人次','20.8%'],
['国内居民出游','2025 全年','65.22 亿人次','16.2%'],
['国内居民出游花费','2025 全年','6.30 万亿元','9.5%'],
['国内出游与总花费','2026 春节 9 天','5.96 亿人次；8034.83 亿元','—'],
['国内出游与总花费','2026 五一假期','3.25 亿人次；1854.92 亿元','—'],
],[4.2,3.15,6.1,3.15],10.5)
note('资料来源：文化和旅游部 [1]—[4]。人次不等于独立用户；国内旅游消费不作为 FlightOR 收入或国际自由行市场金额。春节数据覆盖 9 天。')
table('目标市场的三层界定',['层级','范围','本阶段衡量方式'],[
['TAM','需要数字化搜索、规划和预订的自由行决策需求','说明行业基础，不等同于可获得收入'],
['SAM','中国用户中涉及机票、尤其跨城或跨境航线，且有日期、目的地或多城优化需求的人群','以行为特征界定，暂不套用无统一口径的人口比例'],
['SOM','上线后 12—18 个月可触达的有效规划用户','跟踪有效规划、路线确认、分享、预订跳转与复访'],
],[1.7,8.1,6.8],10.5)
para('主流旅行平台已将 AI 规划与航班搜索结合。Trip.com 提供 TripGenie 和 Trip.Planner，Skyscanner 持续提供灵活日期与目的地探索，Google Flights 也引入 AI 航班优惠搜索。FlightOR 将差异化落实在具体路线决策质量及持续修改体验上。[5]—[9]')

# Competition landscape: consolidate the former competition and differentiation chapters.
section('七  竞争格局')
para('旅行规划领域的主要方案包括 OTA 与 AI 旅行规划、航班搜索工具、旅行内容平台及通用 AI。成熟平台已经覆盖灵活日期、目的地探索和行程生成，FlightOR 聚焦其中对预算、时间和路线组合更敏感的自由行决策场景。')
table('主要竞争对手与同类解决方案',['类别','代表方案','主要能力与价值'],[
['OTA 与 AI 旅行规划','携程、Trip.com（TripGenie、Trip.Planner）、飞猪','连接旅行商品、库存与预订；TripGenie 支持多目的地与协作编辑，Trip.Planner 整合行程及预订'],
['航班搜索与比价','Skyscanner、Google Flights','提供灵活日期、目的地探索与多城市航班搜索；Google Flight Deals 支持自然语言寻找航班优惠'],
['旅行内容平台','小红书、马蜂窝、穷游','提供旅行经验、目的地灵感和避坑信息，帮助用户形成出行意向'],
['通用 AI','ChatGPT 等','理解自然语言、整理信息并生成建议；实时库存能力取决于接入的工具与数据'],
],[3.3,5.2,8.1],10.5)
note('资料来源：[5]—[10]。TripGenie、Trip.Planner 均属于 Trip.com 产品体系；Google Flight Deals 是航班搜索功能。具体可用范围以官方说明的地区、语言和版本为准。')
heading('市场定位')
para('FlightOR 面向预算与时间有限、目的地或路线尚未完全确定的年轻自由行用户，定位为微信场景中的国际航线与路线决策工具。产品连接旅行灵感、航班比较和后续规划，在用户形成路线意向后对接成熟交易平台，早期重点投入预订前的决策效率。')
heading('核心差异化优势')
para('**以航线和约束组织决策。** 围绕预算、日期、兴趣、城市与停留时间组织候选方案，把价格与时间取舍作为规划基础；多城组合、中转利用和路线评分是持续深化的重点。')
para('**让规划状态持续可修改。** 将自然语言需求沉淀为结构化旅行状态，在同一流程中延续条件修改、路线比较和攻略生成，减少用户在内容平台、查票工具与对话之间重复整理信息的工作。')
para('**将生成建议与事实核对分层。** 由 AI 解释与组织方案，通过工具查询航班与价格，在确认阶段重新核价，并保留攻略来源，为用户提供检查和修正方案的依据。')
para('上述差异化主要体现在产品设计与技术架构。微信小程序承载低门槛交互，路线分享与同行讨论继续完善；长期优势将通过真实路线数据、用户修改与选择行为、偏好模型及供应链合作逐步积累，并以真实使用验证规划质量与复访价值。[11]')

# 11 Commercial
section('八  商业模式','收入路径  /  BUSINESS MODEL')
lead('商业化先验证规划价值，再连接交易并拓展高级能力。首条拟验证路径为路线确认后的预订导流。')
diagram('business',['完成路线规划','用户确认路线','查看价格与商品','跳转 OTA 或航司','形成有效成交','按协议结算 CPS'],'首条商业转化路径 合作、转化与收入均待验证',3)
table('分阶段商业化安排',['阶段','方式','范围与进入条件'],[
['1','CPS 与联盟导流','在用户确认航线、酒店或活动后连接合作平台；按实际合作规则计佣'],
['2','会员增值','基础规划保持免费；稳定复访后测试并行路线比较、复杂多城、更高频刷新、价格提醒、协作与跨设备档案'],
['3','酒店与活动等交叉 CPS','围绕已确认行程连接住宿、活动、门票、租车及保险需求'],
['4','B 端 API 或 SaaS','路线引擎稳定后，根据需求探索旅行社、定制游工作室、内容创作者和企业差旅'],
],[1.3,3.8,11.5],10.5,center_cols=(0,))
note('CPS 指按有效成交结算佣金。商业合作、结算费率、成本覆盖和付费意愿均需实测，不预设收入或高付费率。')
para('**闭环验证指标**  路线确认率、报价查看率、预订跳转率、CPS 成交转化率、单次有效规划成本和单个有效规划用户收入。')
para('**推荐原则**  路线排序与商业合作分离，明确标记合作跳转，以用户已确认的旅行目标为依据。')

# 12 Growth
section('九  用户增长与运营策略','用户获取  /  GO TO MARKET')
lead('冷启动阶段聚焦 300—1000 名有真实旅行规划需求的种子用户，建立 20—30 人的持续深访样本池。以上均为招募目标。')
para('首轮以团队现有社交网络和高校旅行社群作为招募起点，重点触达留学生与年轻职场用户；自由行微信群、小红书和 B 站内容受众作为后续补充渠道。冷启动阶段不进行大规模付费投放。')
diagram('growth',['高匹配社群招募','具体旅行问题','首次有效规划','分享与同行讨论','跟踪真实选择','邀请下一批用户'],'冷启动增长与反馈路径 计划',3)
table('内容展示围绕具体路线取舍',['内容示例','展示重点'],[
['北京出发 7 天，预算 6000 元，目的地不限，能去哪','预算、时间和目的地的联合选择'],
['直飞贵 1800 元，中转首尔 16 小时，是否值得','价格节省与时间、过境条件的综合取舍'],
['东京进大阪出与大阪往返，哪个组合更省','城市顺序与进出机场的整体比较'],
],[10.2,6.4])
note('上述金额、时长与城市均为内容选题示例，非实时价格、已验证节省额或用户成交案例。')
para('路线卡分享将展示总价、总时长、城市顺序、核心亮点和对应代价，方便同行人共同讨论。深访围绕真实旅行全过程，记录犹豫节点、信任问题、二次核价和最终选择，并用于改进产品。')

# 13 Technology
section('十  技术与数据能力','能力支撑  /  TECHNOLOGY')
lead('Agent、确定性规则、外部数据与持续维护的产品状态共同支撑旅行规划。')
diagram('architecture',['小程序交互与展示','旅行状态\n条件与方案','Agent 理解与编排','规则与外部工具'],'主要职责关系示意 不表示每次请求均按单向固定流水线执行',4)
table('职责分层',['层次','承担的职责','与用户体验的关系'],[
['Agent','理解需求，选择工具，解释方案并处理多轮修改','降低条件表达和重新整理的负担'],
['确定性规则','预算解析、关键字段约束、缓存及部分降级','减少可避免的解析和约束错误'],
['外部数据与工具','获取航班、搜索、攻略等信息并提供引用','支持价格核对与事实追溯'],
['产品状态与前端','维护旅行条件和方案，呈现路线卡、攻略及操作状态','让用户检查理解、采用结果并继续修改'],
],[3.3,7.5,5.8],10.5)
para('后端统一接入模型与数据供应商，保留适配、异常提示及部分降级机制。关键价格在确认阶段重新查询，服务替换与异常处理尽量减少对前端交互的影响。')
heading('下一阶段的三个技术指标')
para('**规划成功率**、**首个有用结果耗时**、**路线价格与约束准确率**。团队将围绕这三个指标持续评估规划质量与响应效率，改进用户能够实际采用的方案。')
para('在用户授权和隐私合规前提下，逐步积累输入约束、修改或拒绝的方案、最终查看或分享及跳转路线，用于改进路线评分和交互。')
note('技术说明：前端采用 Taro、React 与 MobX；当前对话主链为认证 POST /v1/agent/turns 加短轮询，同步 /v1/agent/converse 共用 Planner。第三方密钥统一在后端管理。详见 [11]。')

# 14 Progress
section('十一  产品进展与验证情况','验证依据  /  PROGRESS AND EVIDENCE')
lead('项目已形成可运行的产品与后端基础，当前处于正式上线前。工程、内容质量、平台体验与市场需求分别验证。')
table('已有产品与工程基础',['产品入口与交互','后端与内容能力'],[
['微信小程序规划入口与自然语言输入','统一 Agent 接口与第三方服务后端管理'],
['结构化旅行条件与多轮补充修改','路线推荐、报价确认与旅行状态保存'],
['路线卡及搜索、探索、我的等页面','按需生成带来源攻略、会话历史、异常提示与部分降级'],
],[8.3,8.3],10.5)
table('历史工程验证记录',['验证范围','已记录结果','证据口径'],[
['根项目自动化测试','140/140','历史记录'],
['后端自动化测试','20 个文件，121 项','历史记录'],
['静态检查与构建','root/backend typecheck、backend build、weapp build、git diff check 通过','历史记录'],
['外部链路','付费模型 Provider 直连、Agent 样例与攻略搜索链路验证','历史样例，不替代完整产品验收'],
],[4,8.9,3.7],10)
note('来源：[11]。上述数量不是本次复测；当前文档仍记录部分内容质量与微信地图等验收缺口，产品构建或结果保存不等于全平台内容可用。')
heading('上线后的有效规划指标')
para('**北极星指标**  每周完成有效路线规划的用户数（Weekly Successful Planners）。')
table('评估维度与观察指标',['维度','观察指标'],[
['激活与质量','首次有效规划比例；成功率、重试率、报价确认成功率、修改后继续规划比例'],
['效率与价值','首个有用结果耗时、交互轮数；收藏、分享、报价确认和预订跳转率'],
['留存与信任','30 天内因另一趟旅行再次使用的比例；价格、签证和时刻等错误反馈比例'],
],[3,13.6],10)

# 15 Roadmap
section('十二  发展规划','验证路径  /  ROADMAP')
lead('未来十二个月围绕真实使用、搜索效率和下一次旅行复访推进。阶段以正式启动上线与种子验证为起点。')
table('十二个月发展安排',['时间','阶段重点','主要工作','验证方向'],[
['0—2 个月','上线与种子验证','正式部署、监控、隐私政策与必要的平台合规；招募 300—1000 名种子用户；修复失败、长等待和错误路线；建立数据看板并每周复盘','用户是否愿意完成真实规划'],
['3—5 个月','深化核心场景','优化灵活日期、低价航线、多城及中转；完善路线评分与分享；记录拒绝和选择原因；重点强化 2—3 个高频出境区域','路线是否减少重复搜索与比较'],
['6—9 个月','留存与交易连接','接入机票或酒店预订跳转及联盟合作；测试价格提醒；根据使用、留存和复访数据确定会员能力','用户是否再次使用并进入预订'],
['9—12 个月','商业化验证','评估 CPS 转化、单个有效规划用户价值；小范围测试高级会员；依据需求判断 B 端能力输出','收入与规划成本能否形成可持续关系'],
],[2.25,3.15,8.1,3.1],10.5)
note('时间区间沿用原计划，6—9 与 9—12 个月在第 9 个月衔接。计划目标不代表已实现运营数据。')
heading('年度核心验证')
para('**真实使用**  用户愿意以实际旅行任务检验 FlightOR。')
para('**决策效率**  产品能够降低跨平台搜索和路线比较成本。')
para('**持续复访**  用户在下一次旅行中再次使用 FlightOR。')

# Core team: confirmed education and internships; no inferred dates, majors or duties.
section('十三  核心团队背景')
para('核心团队由胡斌、曾郡玫和李昊泽组成，教育背景涉及北京航空航天大学、南京大学和清华大学。胡斌与曾郡玫具有本科及硕士教育经历，并曾在多家企业或业务团队实习。')
heading('胡斌')
p=para('**教育背景：**本科阶段就读于北京航空航天大学，硕士阶段就读于清华大学。')
p.paragraph_format.keep_with_next=True
para('**实习经历：**曾在鹰角、百度、高德、可灵等企业或业务团队实习。')
heading('曾郡玫')
p=para('**教育背景：**本科阶段就读于南京大学，硕士阶段就读于清华大学。')
p.paragraph_format.keep_with_next=True
para('**实习经历：**曾在网易、面壁智能和微软实习。')
heading('李昊泽')
para('**教育背景：**目前为北京航空航天大学本科在读学生，是 FlightOR 核心团队成员之一。')

section('十四  风险与应对','经营风险  /  RISK RESPONSE')
table('关键风险与应对安排',['风险','影响','应对安排'],[
['航班数据准确性与可用性','第三方数据可能延迟、缺失或价格变化','规划价与确认价分层，显示数据时间；关键报价二次确认；保留多供应商替换能力'],
['第三方 API 成本','模型多轮调用、航班与网页搜索随用户增长增加成本','缓存重复查询，简单字段采用确定性解析，减少无价值轮次；跟踪单次成功规划成本并按价值选用模型'],
['Agent 稳定性与响应时间','复杂任务调用多，失败或等待过长影响转化','采用超时、重试、阶段性结果与失败提示；持续完善可恢复任务状态、降级和失败案例库'],
['用户信任','价格、时间、签证等重大错误可能破坏信任','标注来源与更新时间，区分生成内容与核实事实；对签证、过境、行李等信息提供官方核验渠道，提示出行前确认'],
['大平台快速复制','OTA 与通用 AI 持续扩展旅行规划','聚焦复杂航线、多城与中转判断；从真实选择与修改中改进评分和交互'],
['商业化与推荐中立性','佣金激励可能影响路线选择','路线排序与商业合作分离，明确标记合作跳转，以用户目标和约束为依据'],
['冷启动与复用不足','用户觉得有趣但没有实际出行需求，难以复用','围绕真实日期招募，以有效规划衡量增长，持续跟踪最终路线采用和后续旅行复访'],
],[3.3,4.9,8.4],10.5)
para('以上安排覆盖既有机制与后续改进。团队将结合实际运行记录和用户反馈逐项验证效果，不使用缺乏依据的风险概率或等级评分。')

# 18 Conclusion and incubation resources
section('十五  阶段目标与结语','孵化目标  /  NEXT STAGE')
lead('FlightOR 已形成可运行的小程序、独立后端与 Agent 规划链路，下一阶段重点验证真实旅行决策中的使用价值。')
para('面对预算有限、日期灵活、目的地或路线尚未完全确定的需求，项目将持续检验：用户能否更快形成愿意采用的路线，是否减少跨平台反复搜索，以及是否在下一次旅行时再次使用。')
table('孵化期的三项核心验证',['验证主题','对应问题','观察方式'],[
['用户价值','是否愿意用 FlightOR 规划真实旅行','记录完整任务、修改过程与最终采用方案'],
['效率价值','是否减少路线搜索与比较工作','观察有用结果耗时、交互轮数及跨平台复核行为'],
['复访价值','下一次旅行是否再次使用','跟踪复访及新旅行规划行为'],
],[3.1,6.5,7])
heading('孵化资源需求')
para('孵化期重点寻求真实用户测试与校园社群资源、创业与运营辅导，以及旅行行业和预订平台对接支持，服务于种子验证、交易路径测试和后续商业化。资金与资源规模依据实际验证进展确定。')
diagram('next_steps',['规划准确与高效','种子用户验证','预订跳转与提醒','高级规划与服务'],'未来一年推进路径',4)
para('未来一年将集中改善航线规划的准确性和响应效率，服务具有真实需求的种子用户，再逐步验证预订跳转、价格提醒与高级规划。完成用户价值与复访验证后，项目将具备拓展完整旅行服务的依据。')

# 19 References
section('主要公开资料来源','参考资料  /  REFERENCES')
note('网络资料核查与访问日期统一为 2026-09-29。数据引用保留原统计年份和口径；产品功能以官方说明对应的地区、语言和版本为准。')
refs=[
('文化和旅游部','2025 年国内居民出游数据情况','2026-01-26','https://zwgk.mct.gov.cn/zfxxgkml/tjxx/202601/t20260126_964367.html'),
('文化和旅游部','2025 年文化和旅游发展统计公报','2026-06-02','https://zwgk.mct.gov.cn/zfxxgkml/tjxx/202606/t20260602_966073.html'),
('文化和旅游部','2026 年春节假期国内出游 5.96 亿人次','2026-02-24','https://www.mct.gov.cn/whzx/whyw/202602/t20260224_964790.htm'),
('文化和旅游部','2026 年五一假期国内出游 3.25 亿人次','2026-05-07','https://www.mct.gov.cn/wlbphone/wlbydd/xxfb/jiaodianxinwen/202605/t20260506_965708.html'),
('Trip.com','Trip.com Launches Trip.Planner: Smart Itineraries Tailored to Your Travel Style with Real-Time Recommendations','2025-08-26','https://www.trip.com/newsroom/trip-com-launches-trip-planner-smart-itineraries-tailored-to-your-travel-style-with-real-time-recommendations/'),
('Trip.com Group','Plan Less, Travel More: TripGenie Unveils New Features for Effortless Travel Planning','2024-06-06','https://www.trip.com/newsroom/tripgenie-new-features-2/'),
('Skyscanner','How do I find the best prices?','更新于 2026-03-26','https://help.skyscanner.net/hc/en-us/articles/201150942-How-do-I-find-the-best-prices'),
('Google Travel Help','How to find the best fares with Google Flights','页面未标注发布日期','https://support.google.com/travel/answer/7664728?hl=en'),
('Google Travel Help','Find flight deals with AI in Google Flights','页面未标注发布日期','https://support.google.com/travel/answer/16497283?hl=en'),
('Alibaba Group','Fliggy','页面未标注发布日期','https://www.alibabagroup.com/en-US/about-alibaba-businesses-1752789467896217600'),
]
def hyperlink(p,label,url):
    rid=p.part.relate_to(url,RT.HYPERLINK,is_external=True);el=OxmlElement('w:hyperlink');el.set(qn('r:id'),rid)
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),BLUE);pr.append(c)
    sz=OxmlElement('w:sz');sz.set(qn('w:val'),'18');pr.append(sz);r.append(pr);txt=OxmlElement('w:t');txt.text=label;r.append(txt);el.append(r);p._p.append(el)
for i,(org,title,date,url) in enumerate(refs,1):
    p=para(f'[{i}] {org}．{title}[EB/OL]．{date}．','BP Ref');hyperlink(p,'官方原文',url)
    if i==7:
        p.add_run('；');hyperlink(p,'多城市搜索说明','https://help.skyscanner.net/hc/en-us/articles/201151512-How-do-I-search-for-multiple-destinations')
    if i==8:
        p.add_run('；');hyperlink(p,'往返及多城市搜索','https://support.google.com/travel/answer/2475306?hl=en')
p=para('[11] FlightOR 项目内部资料．README；docs/PROJECT_CONTEXT.md；docs/archive/legacy/PROJECT_CONTEXT.md；相关实现与验收文档[Z]．核查于 2026-09-29．','BP Ref')
p=para('[12] Intforce．Takeshita Street in December 2018[图像]．Wikimedia Commons．CC BY-SA 4.0．景点详情截图保留原页面作者与许可标识．','BP Ref')
hyperlink(p,'原文件与许可','https://commons.wikimedia.org/wiki/File:Takeshita_Street_in_December_2018.jpg')
p.add_run('；');hyperlink(p,'CC BY-SA 4.0','https://creativecommons.org/licenses/by-sa/4.0/')
note('截图内部记录：FLIGHT_FIRST_ACCEPTANCE.md（2026-09-14）；PUBLICATION_UI_2026-09-22.md；PLACE_MEDIA_2026-09-22.md。航班价格为历史查询；固定测试行程不代表真实用户案例。')
note('内部资料用于说明技术实现及历史工程验证，不作为市场规模、营收或真实用户采用率的依据。')

doc.core_properties.title='FlightOR 商业计划书'
doc.core_properties.subject='创业孵化与上线前验证规划'
doc.core_properties.author='FlightOR 项目团队'
doc.core_properties.keywords='FlightOR,商业计划书,国际自由行,路线规划'
doc.core_properties.comments=''
doc.save(OUT)
# Keep a human-readable source of the final business-plan content under docs.
md=['# FlightOR 商业计划书内容底稿','', '版本：2026-09-29。状态：与当前终稿 DOCX 同步；合并竞争格局，并按用户提供的信息补充核心团队简历。', '']
for el in doc.element.body:
    if el.tag==qn('w:p'):
        txt=''.join(el.xpath('.//w:t/text()'))
        if txt: md += [txt,'']
    elif el.tag==qn('w:tbl'):
        rows=[[''.join(c.xpath('.//w:t/text()')).replace('|','／').replace('\n','<br>') for c in r.xpath('./w:tc')] for r in el.xpath('./w:tr')]
        md+=['| '+' | '.join(rows[0])+' |','| '+' | '.join(['---']*len(rows[0]))+' |']
        md+=['| '+' | '.join(row)+' |' for row in rows[1:]];md+=['']
md+=['## 图表流程文本','']
for pic in doc.inline_shapes:
    alt=pic._inline.docPr.get('descr')
    if alt: md += [alt,'']
md+=['## 网络参考资料链接','']
for i,(_,title,_,url) in enumerate(refs,1): md += [f'{i}. [{title}]({url})']
(ROOT.parent.parent/'docs'/'business-plan'/'FlightOR_BP_2026-09-29.md').write_text('\n'.join(line.rstrip() for block in md for line in block.split('\n'))+'\n',encoding='utf-8')
(ROOT/'qa'/'manifest.json').write_text(json.dumps({'source':str(SOURCE),'source_sha256':hashlib.sha256(SOURCE.read_bytes()).hexdigest(),'output':str(OUT),'sections':sections,'layout':'continuous_report','tables':tblcount,'figures':figcount,'team_section':'user_supplied_brief_bios','sources':refs},ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'docx':str(OUT),'tables':tblcount,'figures':figcount,'logical_sections':len(sections),'size':OUT.stat().st_size},ensure_ascii=False))
